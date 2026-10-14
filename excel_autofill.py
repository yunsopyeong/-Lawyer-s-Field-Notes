from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from financial_parser import FinancialRecord
from excel_template_analyzer import TemplateAnalysis


def autofill_financials(
    template_path: Path,
    output_xlsx_path: Path,
    records: list[FinancialRecord],
    analysis: TemplateAnalysis,
    unmatched_log_path: Path,
) -> None:
    wb = load_workbook(template_path)

    unmatched = []
    by_year_account = {(r.year, r.account_name): r.amount for r in records}

    for ws in wb.worksheets:
        header_year_col = {}
        account_col = None

        for row in range(1, min(ws.max_row, 20) + 1):
            for col in range(1, ws.max_column + 1):
                val = ws.cell(row=row, column=col).value
                if isinstance(val, str) and val.strip().isdigit() and len(val.strip()) == 4:
                    header_year_col[int(val.strip())] = col
                if val == "계정과목명":
                    account_col = col

        if not header_year_col or account_col is None:
            continue

        for row in range(2, ws.max_row + 1):
            account = ws.cell(row=row, column=account_col).value
            if not account or not isinstance(account, str):
                continue
            for year, col in header_year_col.items():
                key = (year, account.strip())
                if key not in by_year_account:
                    continue
                if (ws.title, row, col) in analysis.formula_cells:
                    continue
                target_cell = ws.cell(row=row, column=col)
                if isinstance(target_cell, MergedCell):
                    continue
                target_cell.value = by_year_account[key]

    used_keys = {(r.year, r.account_name) for r in records}
    matched = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for val in row:
                if isinstance(val, str):
                    for y, acc in used_keys:
                        if acc == val:
                            matched.add((y, acc))
    for key in sorted(used_keys - matched):
        unmatched.append(f"미매칭: {key[0]} / {key[1]}")

    unmatched_log_path.parent.mkdir(parents=True, exist_ok=True)
    unmatched_log_path.write_text("\n".join(unmatched), encoding="utf-8")

    output_xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx_path)
