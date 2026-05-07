from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from openpyxl import load_workbook


@dataclass
class TemplateAnalysis:
    formula_cells: set[tuple[str, int, int]]
    merged_top_left_cells: set[tuple[str, int, int]]


def analyze_template(template_path: Path) -> TemplateAnalysis:
    wb = load_workbook(template_path)
    formula_cells: set[tuple[str, int, int]] = set()
    merged_top_left_cells: set[tuple[str, int, int]] = set()

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.add((ws.title, cell.row, cell.column))

        for rng in ws.merged_cells.ranges:
            merged_top_left_cells.add((ws.title, rng.min_row, rng.min_col))

    return TemplateAnalysis(formula_cells=formula_cells, merged_top_left_cells=merged_top_left_cells)
