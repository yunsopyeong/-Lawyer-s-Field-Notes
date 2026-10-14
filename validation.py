from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook

ERROR_TOKENS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "값1", "값2", "폴더경로", "네임"}


def validate_excel_errors(xlsx_path: Path, report_path: Path) -> list[str]:
    wb = load_workbook(xlsx_path, data_only=False)
    findings = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value in ERROR_TOKENS:
                    findings.append(f"{ws.title}!{cell.coordinate} -> {cell.value}")

    report_path.parent.mkdir(parents=True, exist_ok=True)
    if findings:
        report_path.write_text("\n".join(findings), encoding="utf-8")
    else:
        report_path.write_text("오류 없음", encoding="utf-8")
    return findings
